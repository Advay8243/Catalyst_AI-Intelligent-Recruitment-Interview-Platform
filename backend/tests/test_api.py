from io import BytesIO
import secrets

import pytest
from docx import Document

from backend.dependencies import get_email_provider
from backend.main import app
from backend.services.providers import EmailProvider


JOB = {
    "title": "Backend Engineer",
    "company": "Example Labs",
    "description": (
        "Backend Engineer\nBuild SaaS APIs with Python, FastAPI, SQLAlchemy and PostgreSQL. "
        "5+ years experience. Bachelor degree required."
    ),
}


def docx_resume(name: str = "Alex Morgan", email: str = "alex@example.com") -> bytes:
    document = Document()
    document.add_paragraph(name)
    document.add_paragraph(email)
    document.add_paragraph("+1 202 555 0147")
    document.add_paragraph("6 years experience")
    document.add_paragraph("Bachelor")
    document.add_paragraph("SaaS Python FastAPI SQLAlchemy PostgreSQL Docker")
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def create_job(client):
    response = client.post("/jobs", json=JOB)
    assert response.status_code == 201, response.text
    return response.json()


def upload_candidate(client, job_id: str):
    response = client.post(
        f"/jobs/{job_id}/resume",
        files={
            "file": (
                "alex.docx",
                docx_resume(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert response.status_code == 201, response.text
    return response.json()


def test_health_and_job_crud(client):
    assert client.get("/health").json() == {"status": "ok"}
    created = create_job(client)
    assert created["requirements"]["structured_data"]["required_skills"] == [
        "python",
        "fastapi",
        "sqlalchemy",
        "postgresql",
    ]
    assert client.get("/jobs").json()[0]["id"] == created["id"]
    assert client.get(f"/jobs/{created['id']}").json()["title"] == "Backend Engineer"
    assert client.get("/jobs/00000000-0000-0000-0000-000000000000").status_code == 404


def test_multipart_job_description(client):
    response = client.post(
        "/jobs",
        data={"title": "API Developer", "jd_text": "API Developer needs Python and FastAPI. 3 years."},
    )
    assert response.status_code == 201
    assert response.json()["title"] == "API Developer"


def test_resume_upload_candidate_listing_and_analysis(client):
    job = create_job(client)
    uploaded = upload_candidate(client, job["id"])
    candidate = uploaded["candidate"]
    analysis = uploaded["analysis"]
    assert analysis["overall_score"] >= 80
    assert analysis["scoring"]["skills"]["weight"] == 40
    assert "protected traits are excluded" in analysis["explanation"]

    listing = client.get(
        f"/jobs/{job['id']}/candidates",
        params={"search": "alex", "min_score": 50, "sort": "score", "order": "desc"},
    )
    assert listing.status_code == 200, listing.text
    assert listing.json()["total"] == 1
    assert listing.json()["items"][0]["application_status"] == "new"
    assert client.get(f"/candidates/{candidate['id']}").status_code == 200
    detail = client.get(
        f"/candidates/{candidate['id']}/resume-analysis",
        params={"job_id": job["id"]},
    )
    assert detail.status_code == 200
    assert detail.json()["evidence"]["skills"]


def test_resume_upload_infers_docx_when_browser_sends_octet_stream(client):
    job = create_job(client)
    response = client.post(
        f"/jobs/{job['id']}/resume",
        files={"file": ("alex.docx", docx_resume(), "application/octet-stream")},
    )
    assert response.status_code == 201, response.text
    assert response.json()["analysis"]["overall_score"] >= 80


def test_upload_validation_and_duplicate_application(client):
    job = create_job(client)
    bad = client.post(
        f"/jobs/{job['id']}/resume",
        files={"file": ("resume.txt", b"Alex\nalex@example.com", "text/plain")},
    )
    assert bad.status_code == 400
    upload_candidate(client, job["id"])
    duplicate = client.post(
        f"/jobs/{job['id']}/resume",
        files={
            "file": (
                "alex.docx",
                docx_resume(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert duplicate.status_code == 409


def test_missing_resources(client):
    missing = "00000000-0000-0000-0000-000000000000"
    assert client.get(f"/jobs/{missing}/candidates").status_code == 404
    assert client.get(f"/candidates/{missing}").status_code == 404
    assert client.get(f"/candidates/{missing}/resume-analysis").status_code == 404


def test_hr_mutation_requires_authorization(client):
    missing = "00000000-0000-0000-0000-000000000000"
    response = client.post(
        f"/candidates/{missing}/decision",
        json={"job_id": missing, "decision": "advanced"},
        headers={
            "Authorization": f"Bearer {secrets.token_urlsafe(24)}",
            "X-HR-User": "Unauthorized User",
        },
    )
    assert response.status_code == 401


def test_hr_call_transcript_analysis_updates_candidate_score(client):
    job = create_job(client)
    uploaded = upload_candidate(client, job["id"])
    candidate_id = uploaded["candidate"]["id"]

    created = client.post(
        f"/candidates/{candidate_id}/call-sessions",
        json={"job_id": job["id"]},
    )
    assert created.status_code == 201, created.text
    session = created.json()
    assert session["status"] == "not_started"
    assert session["candidate"]["jd_resume_score"] >= 80

    started = client.post(f"/call-sessions/{session['id']}/start")
    assert started.status_code == 200, started.text
    assert started.json()["status"] == "connected"

    entries = [
        ("hr", "Please summarize your relevant experience."),
        (
            "candidate",
            "I have six years of experience and led Python, FastAPI, SQLAlchemy and PostgreSQL projects.",
        ),
        ("hr", "What interests you about this role?"),
        (
            "candidate",
            "I am excited about the role and can start after two weeks notice.",
        ),
    ]
    for speaker, text in entries:
        response = client.post(
            f"/call-sessions/{session['id']}/transcript",
            json={"speaker": speaker, "text": text},
        )
        assert response.status_code == 201, response.text

    completed = client.post(f"/call-sessions/{session['id']}/complete")
    assert completed.status_code == 200, completed.text
    result = completed.json()
    assert result["session"]["status"] == "completed"
    assert result["analysis"]["overall_score"] > 0
    assert result["analysis"]["question_analysis"]

    listing = client.get(f"/jobs/{job['id']}/candidates")
    assert listing.status_code == 200, listing.text
    item = listing.json()["items"][0]
    assert item["hr_score"] == result["analysis"]["overall_score"]
    assert item["call_status"] == "completed"


def complete_screening(client, job_id: str, candidate_id: str):
    created = client.post(
        f"/candidates/{candidate_id}/call-sessions",
        json={"job_id": job_id},
    )
    session_id = created.json()["id"]
    assert client.post(f"/call-sessions/{session_id}/start").status_code == 200
    assert client.post(
        f"/call-sessions/{session_id}/transcript",
        json={
            "speaker": "candidate",
            "text": "I led Python and PostgreSQL services for six years and enjoy collaborative teams.",
        },
    ).status_code == 201
    completed = client.post(f"/call-sessions/{session_id}/complete")
    assert completed.status_code == 200
    return completed.json()


@pytest.mark.parametrize("decision", ["accepted", "advanced", "rejected"])
def test_candidate_review_decision_email_and_history(client, decision):
    job = create_job(client)
    uploaded = upload_candidate(client, job["id"])
    candidate_id = uploaded["candidate"]["id"]
    complete_screening(client, job["id"], candidate_id)

    detail = client.get(
        f"/candidates/{candidate_id}/review", params={"job_id": job["id"]}
    )
    assert detail.status_code == 200, detail.text
    review = detail.json()
    assert review["application"]["screening_status"] == "Awaiting HR Decision"
    assert review["screening_analysis"]["label"] == "AI-assisted screening assessment"
    assert review["call_session"]["transcript"]
    downloaded = client.get(f"/resumes/{review['resume']['id']}/download")
    assert downloaded.status_code == 200
    assert downloaded.content

    decided = client.post(
        f"/candidates/{candidate_id}/decision",
        json={
            "job_id": job["id"],
            "decision": decision,
            "decision_by": "Demo Recruiter",
        },
    )
    assert decided.status_code == 200, decided.text
    assert decided.json()["decision"] == decision
    assert decided.json()["decision_at"]
    assert decided.json()["decision_by"] == "Test HR User"
    assert decided.json()["previous_status"] == "new"
    assert decided.json()["new_status"] == decision.upper()
    listing = client.get(f"/jobs/{job['id']}/candidates")
    listed = listing.json()["items"][0]
    assert listed["screening_status"] == decision.title()
    assert listed["current_stage"] == decision.title()

    draft = client.post(
        f"/candidates/{candidate_id}/email-draft",
        json={"job_id": job["id"], "email_type": decision},
    )
    assert draft.status_code == 200, draft.text
    message = draft.json()
    assert job["title"] in message["subject"]
    assert message["status"] == "Draft"
    assert "score" not in message["body"].lower()
    if decision in {"accepted", "advanced"}:
        assert "Recruiter contact:" in message["body"]
        assert "Test HR User" in message["body"]

    message["subject"] = f"Edited: {message['subject']}"
    message["body"] += "\n\nEdited by HR."
    sent = client.post(
        f"/candidates/{candidate_id}/emails",
        json={
            "job_id": job["id"],
            "draft_id": message["id"],
            "recipient": "recruiting-alias@example.org",
            "subject": message["subject"],
            "body": message["body"],
        },
    )
    assert sent.status_code == 201, sent.text
    assert sent.json()["status"] == "Sent"
    assert sent.json()["recipient"] == "recruiting-alias@example.org"
    assert sent.json()["subject"].startswith("Edited:")

    history = client.get(
        f"/candidates/{candidate_id}/emails", params={"job_id": job["id"]}
    )
    assert history.status_code == 200
    assert len(history.json()) == 1
    assert history.json()[0]["body"].endswith("Edited by HR.")
    assert history.json()[0]["sent_at"]

    refreshed = client.get(
        f"/candidates/{candidate_id}/review", params={"job_id": job["id"]}
    ).json()
    event_types = [event["event_type"] for event in refreshed["timeline"]]
    assert "candidate_created" in event_types
    assert "resume_uploaded" in event_types
    assert "screening_started" in event_types
    assert "screening_completed" in event_types
    assert "decision_made" in event_types
    assert "email_generated" in event_types
    assert "email_sent" in event_types


class FailingEmailProvider(EmailProvider):
    def send_email(self, recipient: str, subject: str, body: str) -> str:
        raise RuntimeError("Simulated provider failure")


def test_mock_email_failure_is_persisted_and_audited(client):
    job = create_job(client)
    uploaded = upload_candidate(client, job["id"])
    candidate_id = uploaded["candidate"]["id"]
    complete_screening(client, job["id"], candidate_id)
    assert client.post(
        f"/candidates/{candidate_id}/decision",
        json={"job_id": job["id"], "decision": "advanced"},
    ).status_code == 200
    draft = client.post(
        f"/candidates/{candidate_id}/email-draft",
        json={"job_id": job["id"], "email_type": "advanced"},
    ).json()

    app.dependency_overrides[get_email_provider] = lambda: FailingEmailProvider()
    try:
        failed = client.post(
            f"/candidates/{candidate_id}/emails",
            json={
                "job_id": job["id"],
                "draft_id": draft["id"],
                "recipient": draft["recipient"],
                "subject": draft["subject"],
                "body": draft["body"],
            },
        )
    finally:
        app.dependency_overrides.pop(get_email_provider, None)

    assert failed.status_code == 400
    history = client.get(
        f"/candidates/{candidate_id}/emails", params={"job_id": job["id"]}
    ).json()
    assert history[0]["status"] == "Failed"
    assert history[0]["sent_at"] is None
    review = client.get(
        f"/candidates/{candidate_id}/review", params={"job_id": job["id"]}
    ).json()
    assert "email_failed" in {
        event["event_type"] for event in review["audit_events"]
    }


def test_further_review_can_transition_once_to_final_decision(client):
    job = create_job(client)
    uploaded = upload_candidate(client, job["id"])
    candidate_id = uploaded["candidate"]["id"]
    complete_screening(client, job["id"], candidate_id)

    further_review = client.post(
        f"/candidates/{candidate_id}/decision",
        json={
            "job_id": job["id"],
            "decision": "needs_review",
            "decision_notes": "Verify the availability response with the candidate.",
        },
    )
    assert further_review.status_code == 200, further_review.text
    assert further_review.json()["decision"] == "needs_review"
    assert further_review.json()["decision_notes"].startswith("Verify")

    accepted = client.post(
        f"/candidates/{candidate_id}/decision",
        json={
            "job_id": job["id"],
            "decision": "accepted",
            "decision_notes": "Availability confirmed by HR.",
        },
    )
    assert accepted.status_code == 200, accepted.text
    assert accepted.json()["decision"] == "accepted"

    duplicate = client.post(
        f"/candidates/{candidate_id}/decision",
        json={"job_id": job["id"], "decision": "rejected"},
    )
    assert duplicate.status_code == 400
    assert "final human decision" in duplicate.json()["detail"]

    review = client.get(
        f"/candidates/{candidate_id}/review", params={"job_id": job["id"]}
    ).json()
    assert len(review["decision_history"]) == 2
    assert review["application"]["decision_notes"] == "Availability confirmed by HR."
    event_types = {event["event_type"] for event in review["audit_events"]}
    assert "further_review_requested" in event_types
    assert "candidate_accepted" in event_types


def test_internal_notification_and_duplicate_email_protection(client):
    job = create_job(client)
    uploaded = upload_candidate(client, job["id"])
    candidate_id = uploaded["candidate"]["id"]
    complete_screening(client, job["id"], candidate_id)
    assert client.post(
        f"/candidates/{candidate_id}/decision",
        json={"job_id": job["id"], "decision": "accepted"},
    ).status_code == 200

    draft_response = client.post(
        f"/candidates/{candidate_id}/email-draft",
        json={
            "job_id": job["id"],
            "email_type": "internal",
            "recipient": "interviewer@example.org",
        },
    )
    assert draft_response.status_code == 200, draft_response.text
    draft = draft_response.json()
    assert draft["recipient"] == "interviewer@example.org"
    assert "Resume score:" in draft["body"]
    assert "HR screening score:" in draft["body"]

    payload = {
        "job_id": job["id"],
        "draft_id": draft["id"],
        "recipient": draft["recipient"],
        "subject": draft["subject"],
        "body": draft["body"],
    }
    first_send = client.post(f"/candidates/{candidate_id}/emails", json=payload)
    assert first_send.status_code == 201, first_send.text
    assert first_send.json()["status"] == "Sent"

    duplicate_send = client.post(
        f"/candidates/{candidate_id}/emails", json=payload
    )
    assert duplicate_send.status_code == 400
    assert "Only draft emails" in duplicate_send.json()["detail"]

    duplicate_draft = client.post(
        f"/candidates/{candidate_id}/email-draft",
        json={
            "job_id": job["id"],
            "email_type": "internal",
            "recipient": "interviewer@example.org",
        },
    )
    assert duplicate_draft.status_code == 400
    assert "already been sent" in duplicate_draft.json()["detail"]


def test_job_parse_preview_create_edit_lifecycle_and_delete(client):
    preview = client.post(
        "/jobs/parse",
        json={
            "title": "Data Engineer",
            "description": (
                "Data Engineer\nDepartment: Analytics\nLocation: Remote\n"
                "Full-time role. Responsibilities:\n- Build pipelines\n- Own Spark jobs\n"
                "Required: Python SQL Spark\nPreferred: Airflow dbt\n"
                "Bachelor degree. AWS Certified preferred. 4+ years."
            ),
        },
    )
    assert preview.status_code == 200, preview.text
    extracted = preview.json()["extracted"]
    assert extracted["title"] == "Data Engineer"
    assert "python" in extracted["required_skills"]

    created = client.post(
        "/jobs",
        json={
            "title": "Data Engineer",
            "department": "Analytics",
            "location": "Remote",
            "employment_type": "full-time",
            "experience_required": "4+ years",
            "description": preview.json()["description"],
            "required_skills": extracted["required_skills"],
            "preferred_skills": ["airflow", "dbt"],
            "responsibilities": ["Build pipelines", "Own Spark jobs"],
            "education": ["bachelor"],
            "certifications": ["aws certified"],
            "status": "draft",
        },
    )
    assert created.status_code == 201, created.text
    job_id = created.json()["id"]
    assert created.json()["status"] == "draft"
    assert created.json()["department"] == "Analytics"

    published = client.post(f"/jobs/{job_id}/publish")
    assert published.status_code == 200
    assert published.json()["status"] == "published"

    updated = client.patch(
        f"/jobs/{job_id}",
        json={"title": "Senior Data Engineer", "location": "Hybrid"},
    )
    assert updated.status_code == 200
    assert updated.json()["title"] == "Senior Data Engineer"

    archived = client.post(f"/jobs/{job_id}/archive")
    assert archived.status_code == 200
    assert archived.json()["status"] == "archived"

    deleted = client.delete(f"/jobs/{job_id}")
    assert deleted.status_code == 204
    assert client.get(f"/jobs/{job_id}").status_code == 404


def test_batch_resume_upload_duplicate_hash_and_isolated_failures(client):
    job = create_job(client)
    good = docx_resume("Alex Morgan", "alex@example.com")
    other = docx_resume("Jordan Lee", "jordan@example.com")
    batch = client.post(
        f"/jobs/{job['id']}/resumes/batch",
        files=[
            (
                "files",
                (
                    "alex.docx",
                    good,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            ),
            (
                "files",
                (
                    "bad.txt",
                    b"not a resume",
                    "text/plain",
                ),
            ),
            (
                "files",
                (
                    "jordan.docx",
                    other,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            ),
            (
                "files",
                (
                    "alex-copy.docx",
                    good,
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            ),
        ],
    )
    assert batch.status_code == 201, batch.text
    payload = batch.json()
    assert payload["success_count"] == 2
    assert payload["failure_count"] == 1
    assert payload["duplicate_count"] == 1
    statuses = {item["filename"]: item["status"] for item in payload["results"]}
    assert statuses["alex.docx"] == "success"
    assert statuses["jordan.docx"] == "success"
    assert statuses["bad.txt"] == "failed"
    assert statuses["alex-copy.docx"] == "duplicate"
    listing = client.get(f"/jobs/{job['id']}/candidates")
    assert listing.json()["total"] == 2
    breakdown = listing.json()["items"][0]["score_breakdown"]
    assert breakdown is not None
    assert "matched_skills" in breakdown


def test_resume_parse_correction_updates_match_and_audit(client):
    job = create_job(client)
    uploaded = upload_candidate(client, job["id"])
    candidate_id = uploaded["candidate"]["id"]
    corrected = client.post(
        f"/candidates/{candidate_id}/resume-parse-correction",
        json={
            "job_id": job["id"],
            "skills": ["python", "fastapi", "sqlalchemy", "postgresql", "docker"],
            "years_experience": 8,
            "certifications": ["aws certified"],
            "projects": ["Built SaaS billing APIs"],
        },
    )
    assert corrected.status_code == 200, corrected.text
    assert corrected.json()["corrected_by"] == "Test HR User"
    assert corrected.json()["analysis"]["overall_score"] >= uploaded["analysis"]["overall_score"]
    review = client.get(
        f"/candidates/{candidate_id}/review", params={"job_id": job["id"]}
    ).json()
    assert review["resume"]["parse_corrected_by"] == "Test HR User"
    assert "docker" in review["resume"]["parsed_data"]["skills"]
    assert "resume_parse_corrected" in {
        event["event_type"] for event in review["audit_events"]
    }


def test_delete_job_with_applications_archives_instead(client):
    job = create_job(client)
    upload_candidate(client, job["id"])
    deleted = client.delete(f"/jobs/{job['id']}")
    assert deleted.status_code == 204
    remaining = client.get(f"/jobs/{job['id']}")
    assert remaining.status_code == 200
    assert remaining.json()["status"] == "archived"


def test_dashboard_stats_and_activity_use_real_data(client):
    empty = client.get("/dashboard/stats")
    assert empty.status_code == 200
    assert empty.json()["total_candidates"] == 0
    assert empty.json()["active_job_descriptions"] == 0

    job = create_job(client)
    uploaded = upload_candidate(client, job["id"])
    candidate_id = uploaded["candidate"]["id"]
    complete_screening(client, job["id"], candidate_id)
    assert client.post(
        f"/candidates/{candidate_id}/decision",
        json={"job_id": job["id"], "decision": "accepted"},
    ).status_code == 200
    draft = client.post(
        f"/candidates/{candidate_id}/email-draft",
        json={"job_id": job["id"], "email_type": "accepted"},
    ).json()
    assert client.post(
        f"/candidates/{candidate_id}/emails",
        json={
            "job_id": job["id"],
            "draft_id": draft["id"],
            "recipient": draft["recipient"],
            "subject": draft["subject"],
            "body": draft["body"],
        },
    ).status_code == 201

    stats = client.get("/dashboard/stats", params={"job_id": job["id"]})
    assert stats.status_code == 200, stats.text
    payload = stats.json()
    assert payload["total_candidates"] == 1
    assert payload["hr_screened"] == 1
    assert payload["accepted"] == 1
    assert payload["emails_sent"] == 1
    assert payload["average_jd_resume_score"] is not None
    assert payload["average_hr_screening_score"] is not None
    assert payload["candidates_per_job"][0]["candidate_count"] == 1

    activity = client.get("/dashboard/activity", params={"job_id": job["id"]})
    assert activity.status_code == 200
    event_types = {item["event_type"] for item in activity.json()["items"]}
    assert "resume_uploaded" in event_types
    assert "email_sent" in event_types
    assert "candidate_accepted" in event_types


def test_remove_application_allows_resume_reupload(client):
    job = create_job(client)
    uploaded = upload_candidate(client, job["id"])
    candidate_id = uploaded["candidate"]["id"]
    removed = client.delete(
        f"/jobs/{job['id']}/candidates/{candidate_id}/application"
    )
    assert removed.status_code == 204
    listing = client.get(f"/jobs/{job['id']}/candidates")
    assert listing.json()["total"] == 0
    reuploaded = upload_candidate(client, job["id"])
    assert reuploaded["candidate"]["email"] == "alex@example.com"
    assert client.get(f"/jobs/{job['id']}/candidates").json()["total"] == 1


def test_candidate_screening_status_and_score_range_filters(client):
    job = create_job(client)
    uploaded = upload_candidate(client, job["id"])
    pending = client.get(
        f"/jobs/{job['id']}/candidates",
        params={"screening_status": "not_screened", "min_score": 50, "max_score": 100},
    )
    assert pending.status_code == 200
    assert pending.json()["total"] == 1
    complete_screening(client, job["id"], uploaded["candidate"]["id"])
    screened = client.get(
        f"/jobs/{job['id']}/candidates",
        params={"screening_status": "awaiting_hr_decision"},
    )
    assert screened.json()["total"] == 1
    still_pending = client.get(
        f"/jobs/{job['id']}/candidates",
        params={"screening_status": "not_screened"},
    )
    assert still_pending.json()["total"] == 0


def test_jd_summary_bullets_scoring_criteria_and_semantic_search(client):
    created = create_job(client)
    structured = created["requirements"]["structured_data"]
    bullets = structured["summary_bullets"]
    assert 5 <= len(bullets) <= 10
    assert any("python" in bullet.lower() or "fastapi" in bullet.lower() for bullet in bullets)

    criteria = client.get("/scoring-criteria")
    assert criteria.status_code == 200
    payload = criteria.json()
    labels = {item["label"] for item in payload["criteria"]}
    assert "Required Skills" in labels
    assert "Experience" in labels
    assert payload["signals"]

    search = client.get("/jobs/search", params={"q": "backend python fastapi postgresql"})
    assert search.status_code == 200
    hits = search.json()["items"]
    assert hits
    assert hits[0]["job"]["id"] == created["id"]
    assert 0 < hits[0]["similarity"] <= 1


def test_shortlist_threshold_and_generate_scores(client):
    job = create_job(client)
    uploaded = upload_candidate(client, job["id"])
    listing = client.get(f"/jobs/{job['id']}/candidates")
    assert listing.status_code == 200
    assert listing.json()["shortlisted_threshold"] == 60
    assert listing.json()["total_uploaded"] == 1
    assert listing.json()["items"][0]["jd_score"] >= 60
    assert listing.json()["items"][0]["fit_points"]

    generated = client.post(f"/jobs/{job['id']}/generate-scores")
    assert generated.status_code == 200, generated.text
    body = generated.json()
    assert body["analyzed_count"] == 1
    assert body["shortlisted_count"] == 1
    assert body["job_id"] == job["id"]
    assert uploaded["candidate"]["id"]


def test_candidate_search_filters_sort_and_comparison(client):
    job = create_job(client)
    second_job = client.post(
        "/jobs",
        json={
            **JOB,
            "title": "Platform Engineer",
            "description": (
                "Platform Engineer\nBuild platforms with Python and Kubernetes. "
                "4+ years experience. Bachelor degree required."
            ),
        },
    ).json()

    alex = upload_candidate(client, job["id"])
    maya = client.post(
        f"/jobs/{job['id']}/resume",
        files={
            "file": (
                "maya.docx",
                docx_resume(name="Maya Chen", email="maya@example.com"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert maya.status_code == 201, maya.text
    maya_id = maya.json()["candidate"]["id"]
    alex_id = alex["candidate"]["id"]

    by_name = client.get(f"/jobs/{job['id']}/candidates", params={"search": "maya"})
    assert by_name.status_code == 200
    assert by_name.json()["total"] == 1
    assert by_name.json()["items"][0]["full_name"] == "Maya Chen"

    by_email = client.get("/candidates", params={"job_id": job["id"], "search": "alex@example.com"})
    assert by_email.json()["total"] == 1

    by_phone = client.get(f"/jobs/{job['id']}/candidates", params={"search": "555 0147"})
    assert by_phone.json()["total"] >= 1

    by_skill = client.get(f"/jobs/{job['id']}/candidates", params={"search": "fastapi"})
    assert by_skill.json()["total"] >= 1

    by_title = client.get("/candidates", params={"search": "Platform Engineer"})
    assert by_title.status_code == 200
    # No candidates on platform job yet.
    assert by_title.json()["total"] == 0

    client.post(
        f"/jobs/{second_job['id']}/resume",
        files={
            "file": (
                "sam.docx",
                docx_resume(name="Sam Lee", email="sam@example.com"),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    by_title_hit = client.get("/candidates", params={"search": "Platform Engineer"})
    assert by_title_hit.json()["total"] == 1
    assert by_title_hit.json()["items"][0]["job_title"] == "Platform Engineer"

    complete_screening(client, job["id"], alex_id)
    decided = client.post(
        f"/candidates/{alex_id}/decision",
        json={"job_id": job["id"], "decision": "advanced"},
    )
    assert decided.status_code == 200

    draft = client.post(
        f"/candidates/{alex_id}/email-draft",
        json={"job_id": job["id"], "email_type": "advanced"},
    )
    assert draft.status_code == 200
    message = draft.json()
    sent = client.post(
        f"/candidates/{alex_id}/emails",
        json={
            "job_id": job["id"],
            "draft_id": message["id"],
            "recipient": "recruiting-alias@example.org",
            "subject": message["subject"],
            "body": message["body"],
        },
    )
    assert sent.status_code == 201

    decision_filter = client.get(
        f"/jobs/{job['id']}/candidates",
        params={"decision_status": "advanced"},
    )
    assert decision_filter.json()["total"] == 1
    assert decision_filter.json()["items"][0]["decision_status"] == "advanced"

    email_filter = client.get(
        f"/jobs/{job['id']}/candidates",
        params={"email_status": "sent"},
    )
    assert email_filter.json()["total"] == 1

    hr_filter = client.get(
        f"/jobs/{job['id']}/candidates",
        params={"min_hr_score": 1, "max_hr_score": 100},
    )
    assert hr_filter.json()["total"] == 1

    from datetime import date, timedelta

    today = date.today().isoformat()
    tomorrow = (date.today() + timedelta(days=1)).isoformat()
    date_filter = client.get(
        f"/jobs/{job['id']}/candidates",
        params={"uploaded_from": today, "uploaded_to": tomorrow},
    )
    assert date_filter.json()["total"] == 2

    sorted_by_name = client.get(
        f"/jobs/{job['id']}/candidates",
        params={"sort": "name", "order": "asc"},
    )
    names = [item["full_name"] for item in sorted_by_name.json()["items"]]
    assert names == sorted(names)

    sorted_by_decision = client.get(
        f"/jobs/{job['id']}/candidates",
        params={"sort": "decision", "order": "asc"},
    )
    assert sorted_by_decision.status_code == 200
    assert sorted_by_decision.json()["total"] == 2

    page = client.get(
        f"/jobs/{job['id']}/candidates",
        params={"page": 1, "page_size": 1, "sort": "name", "order": "asc"},
    )
    assert page.json()["page_size"] == 1
    assert page.json()["total"] == 2
    assert len(page.json()["items"]) == 1

    compare = client.get(
        f"/jobs/{job['id']}/candidates/compare",
        params=[("ids", alex_id), ("ids", maya_id)],
    )
    assert compare.status_code == 200, compare.text
    payload = compare.json()
    assert payload["job_id"] == job["id"]
    assert len(payload["items"]) == 2
    first = payload["items"][0]
    assert "required_skills_score" in first
    assert "preferred_skills_score" in first
    assert "experience_score" in first
    assert "responsibilities_score" in first
    assert "education_score" in first
    assert "strengths" in first
    assert "missing_information" in first
    assert "ai_recommendation" in first
    assert "human_decision" in first
    assert first["score_breakdown"] is not None or first["jd_score"] is not None
    # Protected characteristics must not appear as comparison fields.
    assert "age" not in first
    assert "gender" not in first
    assert "race" not in first

    too_few = client.get(
        f"/jobs/{job['id']}/candidates/compare",
        params=[("ids", alex_id)],
    )
    assert too_few.status_code == 422
